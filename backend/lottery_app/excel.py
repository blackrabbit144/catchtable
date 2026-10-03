"""응모자 Excel 출력/취입. openpyxl 사용."""
import io

from openpyxl import Workbook, load_workbook

# 출력 컬럼 순서 (1-indexed). 고유번호가 재취입 시 매칭 키.
HEADERS = ['고유번호', '이름', '전화번호', '생년월일', '자녀', '응모일시', '당첨여부']
COL_ENTRY_NO = 1
COL_WINNER = 7

# 당첨여부 칸에 스태프가 기입하면 당첨으로 인정하는 값 (대소문자 무시)
WINNER_MARKS = {'o', 'ㅇ', '당첨', '1', 'y', 'yes', 'true', 'v', 'o', '✓', '○'}


def build_workbook(applicants) -> bytes:
    """응모자 목록을 xlsx 바이트로 반환."""
    wb = Workbook()
    ws = wb.active
    ws.title = '응모자'
    ws.append(HEADERS)
    for a in applicants:
        if a.is_child:
            child = '자녀'
        elif a.children_count > 0:
            child = f'자녀 {a.children_count}명 동반'
        else:
            child = ''
        ws.append([
            a.entry_no,
            a.name,
            a.phone,
            a.birthdate.strftime('%Y-%m-%d') if a.birthdate else '',
            child,
            a.registered_at.strftime('%Y-%m-%d %H:%M') if a.registered_at else '',
            'O' if a.is_winner else '',
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_winner_entry_nos(file_obj) -> list:
    """업로드된 xlsx 에서 당첨여부 칸이 표시된 행의 고유번호 목록을 반환."""
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    winners = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # 헤더
        if not row or len(row) < COL_WINNER:
            continue
        entry_cell = row[COL_ENTRY_NO - 1]
        winner_cell = row[COL_WINNER - 1]
        if entry_cell is None or winner_cell is None:
            continue
        if str(winner_cell).strip().lower() in WINNER_MARKS:
            try:
                winners.append(int(entry_cell))
            except (ValueError, TypeError):
                continue
    return winners
