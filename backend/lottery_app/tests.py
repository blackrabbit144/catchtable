import io
from datetime import date
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from .excel import build_workbook, COL_WINNER
from .models import Store, LotterySettings, LotteryApplicant


class LotteryBase(APITestCase):
    def setUp(self):
        self.store = Store.objects.create(name='부산점', code='busan')
        self.settings = LotterySettings.objects.create(
            store=self.store, is_open=True, registration_token='TOK', max_per_phone=10,
        )

    def _register(self, phone='010-1111-2222', name='홍길동', children_count=0,
                  token='TOK', code='busan', birthdate='1990-01-01'):
        return self.client.post('/api/lottery/register/', {
            'store': code, 'token': token, 'name': name,
            'phone': phone, 'birthdate': birthdate, 'children_count': children_count,
        }, format='json')


class RegisterTests(LotteryBase):
    def test_register_ok_and_receipt_hides_entry_no(self):
        res = self._register()
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        # 고객 응답(명세서)에는 고유번호가 없어야 한다
        self.assertNotIn('entry_no', res.data)
        self.assertIn('public_token', res.data)
        self.assertEqual(res.data['name'], '홍길동')

        token = res.data['public_token']
        receipt = self.client.get(f'/api/lottery/receipt/{token}/')
        self.assertEqual(receipt.status_code, 200)
        self.assertNotIn('entry_no', receipt.data)
        self.assertEqual(receipt.data['store_name'], '부산점')

    def test_closed_rejected(self):
        self.settings.is_open = False
        self.settings.save()
        self.assertEqual(self._register().status_code, status.HTTP_403_FORBIDDEN)

    def test_invalid_token_rejected(self):
        self.assertEqual(self._register(token='WRONG').status_code, status.HTTP_403_FORBIDDEN)

    def test_reregister_replaces_family(self):
        # 1번호=1등록. 같은 번호 재등록은 이전 가족을 교체한다.
        # 자녀 없이 재등록 → 최신만 남음
        self.assertEqual(self._register(name='처음').status_code, 201)
        self.assertEqual(self._register(name='수정').status_code, 201)
        rows = LotteryApplicant.objects.filter(store=self.store, phone='010-1111-2222')
        self.assertEqual(rows.count(), 1)
        self.assertEqual(rows.first().name, '수정')

    def test_reregister_with_children_replaces_not_stacks(self):
        # 자녀2로 등록(3행) → 같은 번호로 자녀1 재등록 → 적재되지 않고 교체(2행)
        self._register(children_count=2)
        self.assertEqual(LotteryApplicant.objects.filter(store=self.store, phone='010-1111-2222').count(), 3)
        self._register(children_count=1)
        rows = LotteryApplicant.objects.filter(store=self.store, phone='010-1111-2222')
        self.assertEqual(rows.count(), 2)
        self.assertEqual(rows.get(is_child=False).children_count, 1)

    def test_children_create_extra_rows_same_phone(self):
        # 자녀 3명 → 부모1 + 자녀3 = 4행 (1회 등록 내)
        res = self._register(children_count=3)
        self.assertEqual(res.status_code, 201)
        rows = LotteryApplicant.objects.filter(store=self.store, phone='010-1111-2222')
        self.assertEqual(rows.count(), 4)
        self.assertEqual(rows.filter(is_child=True).count(), 3)
        parent = rows.get(is_child=False)
        self.assertEqual(parent.children_count, 3)
        # 명세서(고객 응답)에 자녀 수가 보이고 고유번호는 없음
        self.assertEqual(res.data['children_count'], 3)
        self.assertNotIn('entry_no', res.data)

    def test_limit_on_family_size(self):
        # 부모1 + 자녀10 = 11 > max_per_phone(10) → limit
        res = self._register(children_count=10)
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data['detail'], 'limit')

    def test_cancel_removes_whole_family(self):
        res = self._register(children_count=2)
        token = res.data['public_token']
        cancel = self.client.delete(f'/api/lottery/receipt/{token}/cancel/')
        self.assertEqual(cancel.status_code, 200)
        self.assertEqual(LotteryApplicant.objects.filter(store=self.store).count(), 0)

    def test_entry_no_sequential_per_store(self):
        self._register(phone='010-0000-0001')
        self._register(phone='010-0000-0002')
        nos = list(LotteryApplicant.objects.filter(store=self.store).order_by('entry_no').values_list('entry_no', flat=True))
        self.assertEqual(nos, [1, 2])


class StoreIsolationTests(LotteryBase):
    def setUp(self):
        super().setUp()
        self.store2 = Store.objects.create(name='광주점', code='gwangju')
        LotterySettings.objects.create(store=self.store2, is_open=True, registration_token='TOK2')

    def test_same_phone_independent_across_stores(self):
        self.assertEqual(self._register(code='busan', token='TOK').status_code, 201)
        # 別店舗なら同じ電話でも新規として通る
        self.assertEqual(self._register(code='gwangju', token='TOK2').status_code, 201)
        # entry_no は各店舗で 1 から
        self.assertEqual(LotteryApplicant.objects.get(store=self.store).entry_no, 1)
        self.assertEqual(LotteryApplicant.objects.get(store=self.store2).entry_no, 1)

    def test_pickup_toggle(self):
        self._register()
        a = LotteryApplicant.objects.get(store=self.store, entry_no=1)
        # 체크 → picked_up_at 기록
        res = self.client.post(f'/api/lottery/admin/busan/applicant/{a.entry_no}/pickup/',
                               {'picked_up': True}, format='json')
        self.assertEqual(res.status_code, 200)
        a.refresh_from_db()
        self.assertIsNotNone(a.picked_up_at)
        # 체크 해제 → None
        self.client.post(f'/api/lottery/admin/busan/applicant/{a.entry_no}/pickup/',
                         {'picked_up': False}, format='json')
        a.refresh_from_db()
        self.assertIsNone(a.picked_up_at)

    def test_admin_list_scoped_by_store(self):
        self._register(code='busan', token='TOK')
        res = self.client.get('/api/lottery/admin/gwangju/applicants/')
        self.assertEqual(res.data['count'], 0)
        res = self.client.get('/api/lottery/admin/busan/applicants/')
        self.assertEqual(res.data['count'], 1)
        self.assertIn('entry_no', res.data['applicants'][0])  # 管理用は固有番号あり


class ExcelTests(LotteryBase):
    def _make_applicants(self, n=3):
        for i in range(n):
            LotteryApplicant.objects.create(
                store=self.store, entry_no=i + 1, name=f'사람{i+1}',
                phone=f'010-0000-000{i+1}', birthdate=date(1990, 1, 1),
            )

    def _winner_file(self, winner_entry_nos):
        """export 결과를 받아 당첨여부 칸을 채워 업로드용 파일로 반환."""
        content = build_workbook(LotteryApplicant.objects.filter(store=self.store).order_by('entry_no'))
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            entry_no = row[0].value
            if entry_no in winner_entry_nos:
                row[COL_WINNER - 1].value = 'O'
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile(
            'w.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_export_has_headers(self):
        self._make_applicants(2)
        res = self.client.get('/api/lottery/admin/busan/export/')
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(io.BytesIO(res.content))
        self.assertEqual(wb.active.cell(row=1, column=1).value, '고유번호')

    def test_import_marks_winners_only_no_sms(self):
        # 취입은 당첨 표시만 (SMS 발송 없음)
        self._make_applicants(3)
        res = self.client.post(
            '/api/lottery/admin/busan/import/',
            {'file': self._winner_file([1, 3])}, format='multipart',
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['marked'], 2)
        self.assertEqual(res.data['total_winners'], 2)
        self.assertEqual(res.data['unsent'], 2)
        self.assertTrue(LotteryApplicant.objects.get(store=self.store, entry_no=1).is_winner)
        self.assertFalse(LotteryApplicant.objects.get(store=self.store, entry_no=2).is_winner)
        # 아직 아무도 발송되지 않음
        self.assertEqual(
            LotteryApplicant.objects.filter(store=self.store, notified_at__isnull=False).count(), 0)

    @patch('lottery_app.views.send_sms', return_value=True)
    def test_send_sms_substitutes_name_and_store(self, mock_sms):
        self._make_applicants(2)
        self.client.post('/api/lottery/admin/busan/import/',
                         {'file': self._winner_file([1, 2])}, format='multipart')
        res = self.client.post('/api/lottery/admin/busan/send-sms/',
                               {'message': '[{매장}] {이름}님 당첨!'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['sent'], 2)
        # {이름}/{매장} 치환 확인
        texts = [c.args[1] for c in mock_sms.call_args_list]
        self.assertIn('[부산점] 사람1님 당첨!', texts)
        self.assertIn('[부산점] 사람2님 당첨!', texts)

    @patch('lottery_app.views.send_sms', return_value=True)
    def test_send_sms_idempotent_only_unsent(self, mock_sms):
        self._make_applicants(2)
        self.client.post('/api/lottery/admin/busan/import/',
                         {'file': self._winner_file([1, 2])}, format='multipart')
        self.client.post('/api/lottery/admin/busan/send-sms/',
                         {'message': '{이름}'}, format='json')
        # 再送信 → 既発送者には送らない
        res = self.client.post('/api/lottery/admin/busan/send-sms/',
                               {'message': '{이름}'}, format='json')
        self.assertEqual(res.data['target'], 0)
        self.assertEqual(res.data['sent'], 0)
        self.assertEqual(mock_sms.call_count, 2)  # 合計2回のみ

    @patch('lottery_app.views.send_sms', return_value=True)
    def test_send_sms_dedups_by_phone(self, mock_sms):
        # 같은 번호(부모+자녀 3)가 모두 당첨 → SMS 는 1통만
        LotteryApplicant.objects.create(store=self.store, entry_no=1, name='조유진',
                                        phone='010-8888-9999', birthdate=date(1990, 1, 1),
                                        children_count=3, is_winner=True)
        for i in range(3):
            LotteryApplicant.objects.create(store=self.store, entry_no=2 + i,
                                            name=f'조유진 (자녀 {i+1})', phone='010-8888-9999',
                                            birthdate=date(1990, 1, 1), is_child=True, is_winner=True)
        res = self.client.post('/api/lottery/admin/busan/send-sms/',
                               {'message': '당첨자: {이름}'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['target'], 1)       # 번호 1건
        self.assertEqual(res.data['sent'], 1)         # 1통만
        self.assertEqual(res.data['recipients'], 4)   # 당첨자 4명 포함
        self.assertEqual(mock_sms.call_count, 1)      # SMS 호출 1회
        # 4명 모두 발송 완료 처리
        self.assertEqual(
            LotteryApplicant.objects.filter(store=self.store, notified_at__isnull=True).count(), 0)
        # 1통 안에 당첨자 전원 이름이 들어감
        self.assertIn('조유진 (자녀 1)', mock_sms.call_args_list[0].args[1])

    def test_send_sms_empty_message_rejected(self):
        self._make_applicants(1)
        self.client.post('/api/lottery/admin/busan/import/',
                         {'file': self._winner_file([1])}, format='multipart')
        res = self.client.post('/api/lottery/admin/busan/send-sms/',
                               {'message': '   '}, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
